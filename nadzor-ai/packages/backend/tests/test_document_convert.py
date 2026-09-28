"""Приём DOCX и XML (ТЗ 9.1): преобразование в PDF, отказ с причиной.

Проверяется главное: текст исходного документа доезжает до PDF и читается
тем же разбором, что и обычный PDF, — иначе приём формата был бы видимостью.
"""
import hashlib
import uuid
import io
import json
import sys
from pathlib import Path

import pymupdf
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import document_convert, models, official_api  # noqa: E402
from app.db import get_session  # noqa: E402
from app.main import app  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

client = TestClient(app)


def _docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph("Класс бетона В30 по проекту")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Параметр"
    table.rows[0].cells[1].text = "Значение 120 м²"
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


XML = "<?xml version='1.0' encoding='utf-8'?><act><number>42</number>" \
      "<work>Армирование плиты</work></act>".encode()


def _text(pdf: bytes) -> str:
    with pymupdf.open(stream=pdf, filetype="pdf") as doc:
        return " ".join(page.get_text() for page in doc)


def test_formats_are_detected_by_content_not_extension():
    assert document_convert.detect(b"%PDF-1.7 ...") == "PDF"
    assert document_convert.detect(_docx()) == "DOCX"
    assert document_convert.detect(XML) == "XML"
    assert document_convert.detect(b"PK\x03\x04not-a-docx") is None
    assert document_convert.detect(b"\x89PNG....") is None


def test_docx_text_and_tables_reach_the_pdf():
    data = _docx()
    converted = document_convert.to_pdf(data, "акт.docx")
    text = _text(converted.pdf)
    assert "Класс бетона В30" in text and "120 м²" in text
    assert converted.source_format == "DOCX"
    assert converted.source_sha256 == hashlib.sha256(data).hexdigest()


def test_xml_elements_reach_the_pdf_with_their_path():
    converted = document_convert.to_pdf(XML, "акт.xml")
    text = _text(converted.pdf)
    assert "act/work: Армирование плиты" in text
    assert converted.source_sha256 == hashlib.sha256(XML).hexdigest()


def test_xml_with_dtd_is_refused():
    """XXE и раздувание сущностей идут через DTD — такой XML не принимается."""
    bomb = b'<?xml version="1.0"?><!DOCTYPE a [<!ENTITY x "xx">]><a>&x;</a>'
    with pytest.raises(document_convert.UnsupportedFormatError, match="DTD"):
        document_convert.to_pdf(bomb, "a.xml")


def test_broken_docx_is_refused_with_a_reason():
    with pytest.raises(document_convert.UnsupportedFormatError, match="неподдерживаемый"):
        document_convert.to_pdf(b"PK\x03\x04broken", "x.docx")


def test_api_accepts_docx_and_records_the_source(monkeypatch):
    monkeypatch.setattr(official_api, "_execute", lambda run_id: None)
    data = _docx()
    registry = [{"file_id": f"a-{uuid.uuid4().hex[:8]}", "file_name": "act.docx",
                 "object_id": "OBJ-D", "doc_stage": "ID", "discipline": "АР", "document_code": "ID-1", "revision": "1",
                 "approval_status": "APPROVED",
                 "sha256": hashlib.sha256(data).hexdigest()}]
    body = client.post("/api/v1/documents/upload", files=[
        ("files", ("act.docx", data, "application/octet-stream")),
        ("registry", ("r.json", json.dumps(registry).encode(), "application/json")),
    ]).json()
    assert body["rejected"] == [], body
    db = next(get_session())
    try:
        doc = db.get(models.Document, body["accepted"][0]["document_id"])
        assert doc.source_metadata["source_format"] == "DOCX"
        assert doc.source_metadata["source_sha256"] == hashlib.sha256(data).hexdigest()
    finally:
        db.close()
