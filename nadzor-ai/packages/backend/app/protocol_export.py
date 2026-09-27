"""Выгрузка протокола в XML, DOCX и PDF (ТЗ 7, п.7).

Содержимое одно и то же во всех форматах — собранный `protocol.build`:
статус загрузки, сценарий, версии и пять таблиц. Формат меняет только
оформление, поэтому протокол в PDF не может разойтись с протоколом в JSON.
"""
from __future__ import annotations

import io
import xml.etree.ElementTree as ET
from pathlib import Path

_TABLE_TITLES = {
    "completeness": "1. Комплектность и сопоставимость",
    "candidates": "2. Предварительные кандидаты",
    "confirmed_violations": "3. Подтверждённые инспектором нарушения",
    "negative_verified": "4. Проверенные отрицательные результаты",
    "suspicions": "5. Гипотезы свободного поиска",
}
_FONT_CANDIDATES = (
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/dejavu/DejaVuSans.ttf",
)


def _rows(table: str, items: list[dict]) -> list[list[str]]:
    if table == "completeness":
        return [[f'{i.get("parameter_code") or ""} {i.get("parameter_name") or ""}'.strip(),
                 str(i.get("completeness_status") or ""), str(i.get("reason") or "")]
                for i in items]
    rows = []
    for card in items:
        pages = "; ".join(f'{s.get("stage")} {s.get("document_code") or ""} '
                          f'ред.{s.get("revision") or "?"} стр.{s.get("page")}'
                          for s in card.get("sources") or [])
        title = f'{card.get("parameter_code") or ""} {card.get("parameter_name") or ""}'
        rows.append([title.strip(),
                     str(card.get("expected_value") or ""), str(card.get("actual_value") or ""),
                     pages, str(card.get("inspector_decision") or "")])
    return rows


def _header(table: str) -> list[str]:
    if table == "completeness":
        return ["Параметр", "Полнота", "Причина"]
    return ["Параметр", "Ожидается", "Факт", "Источники", "Решение"]


def _summary(payload: dict) -> list[str]:
    proto = payload["protocol"]
    versions = proto["versions"]
    upload = ", ".join(proto["upload_status"].values())
    return [
        f'Процесс: {payload["process_id"]}   Объект: {payload.get("object_id") or ""}',
        f'Статус: {proto["status"]}   Верификация: {proto["verification_status"]}',
        f'Статус загрузки документов: {upload}',
        f'Тип проверки: {proto["scenario"]}',
        f'Версии: матрица {versions["matrix_version"]}, модель {versions["model_version"]}, '
        f'набор данных {versions["dataset_version"]}',
        f'Входной манифест: {versions["input_manifest_hash"]}',
        "Результаты — гипотезы для проверки инспектором, а не заключение о нарушении.",
    ]


def to_xml(payload: dict) -> bytes:
    proto = payload["protocol"]
    root = ET.Element("protocol", process_id=str(payload["process_id"]),
                      object_id=str(payload.get("object_id") or ""),
                      status=proto["status"], verification_status=proto["verification_status"],
                      scenario=proto["scenario"])
    versions = ET.SubElement(root, "versions")
    for key, value in proto["versions"].items():
        versions.set(key, str(value))
    upload = ET.SubElement(root, "upload_status")
    for stage, value in proto["upload_status"].items():
        ET.SubElement(upload, "stage", code=stage, status=value)
    for table, items in proto["tables"].items():
        node = ET.SubElement(root, table)
        for item in items:
            entry = ET.SubElement(node, "item")
            for key, value in item.items():
                if key == "sources":
                    for source in value or []:
                        ET.SubElement(entry, "source", {k: str(v) for k, v in source.items()
                                                        if v is not None})
                elif value is not None:
                    entry.set(key, str(value))
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def to_docx(payload: dict) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Протокол проверки ПД / РД / ИД", level=1)
    for line in _summary(payload):
        document.add_paragraph(line)
    for table, items in payload["protocol"]["tables"].items():
        document.add_heading(_TABLE_TITLES[table], level=2)
        if not items:
            document.add_paragraph("Записей нет.")
            continue
        header = _header(table)
        grid = document.add_table(rows=1, cols=len(header))
        grid.style = "Table Grid"
        for cell, text in zip(grid.rows[0].cells, header, strict=True):
            cell.text = text
        for row in _rows(table, items):
            for cell, text in zip(grid.add_row().cells, row, strict=True):
                cell.text = text
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def _font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for path in _FONT_CANDIDATES:
        if Path(path).is_file():
            pdfmetrics.registerFont(TTFont("ProtocolSans", path))
            return "ProtocolSans"
    # Без шрифта с кириллицей PDF вышел бы набором пустых квадратов —
    # честнее отказать, чем выдать нечитаемый протокол.
    raise RuntimeError("шрифт с кириллицей не найден: установите fonts-dejavu-core")


def to_pdf(payload: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font = _font()
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = font
    stream = io.BytesIO()
    story = [Paragraph("Протокол проверки ПД / РД / ИД", styles["Title"])]
    story += [Paragraph(line, styles["Normal"]) for line in _summary(payload)]
    for table, items in payload["protocol"]["tables"].items():
        story += [Spacer(1, 8), Paragraph(_TABLE_TITLES[table], styles["Heading2"])]
        if not items:
            story.append(Paragraph("Записей нет.", styles["Normal"]))
            continue
        data = [_header(table)] + [[Paragraph(cell, styles["BodyText"]) for cell in row]
                                   for row in _rows(table, items)]
        grid = Table(data, repeatRows=1)
        grid.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), font),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(grid)
    SimpleDocTemplate(stream, pagesize=landscape(A4)).build(story)
    return stream.getvalue()
